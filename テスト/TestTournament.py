"""架空データでTournament.pyの要件とExcel出力形式を検証する。
リポジトリのルートから実行: python テスト/TestTournament.py
認証情報・TSV・既存のExcelを読み書きしない。
"""
from collections import Counter
import contextlib
import io
from pathlib import Path
import random
import sys
import types
import xml.etree.ElementTree as ET
import zipfile

import openpyxl


ROOT = Path(__file__).resolve().parents[1]


def load_module(name, path):
	module = types.ModuleType(name)
	sys.modules[name] = module
	exec(compile(path.read_bytes(), str(path), "exec"), module.__dict__)
	return module


def partitions(total, largest=None):
	if total == 0:
		yield []
		return
	for first in range(min(total, largest or total), 0, -1):
		for rest in partitions(total - first, first):
			yield [first, *rest]


def make_players(module, groups):
	return [
		module.Participant(f"選手{index}", None if index % 2 else f"かな{index}", group)
		for index, group in enumerate(groups)
	]


def check_draw(module, players, seed):
	pairs = module.make_first_round_pairs(players, seed)
	slots = [player for pair in pairs for player in pair]
	real = [player for player in slots if player.name != "BYE"]
	assert Counter(player.name for player in real) == Counter(player.name for player in players)
	assert len(slots) == 1 << (len(players) - 1).bit_length()
	assert all(left.name != "BYE" or right.name != "BYE" for left, right in pairs)

	# 全分割段階で、不戦勝数と各団体の人数差が1以内であること。
	width = len(slots)
	groups = {player.groupname for player in players if player.groupname}
	while width >= 2:
		blocks = [slots[start:start + width] for start in range(0, len(slots), width)]
		byes = [sum(player.name == "BYE" for player in block) for block in blocks]
		assert max(byes) - min(byes) <= 1, (len(players), seed, width, byes)
		for group in groups:
			counts = [sum(player.name != "BYE" and player.groupname == group for player in block) for block in blocks]
			assert max(counts) - min(counts) <= 1, (len(players), seed, width, group, counts)
		width //= 2

	# 両者とも初戦のカードは互いに参加者を共有しない。
	# Kカードなら、最大団体の人数がN-Kを超えた分は同団体対戦が不可避。
	first_meetings = [(left, right) for left, right in pairs if right.name != "BYE"]
	for index in range(0, len(pairs) - 1, 2):
		first, second = pairs[index:index + 2]
		if first[1].name == second[1].name == "BYE":
			first_meetings.append((first[0], second[0]))
	conflicts = sum(bool(left.groupname) and left.groupname == right.groupname for left, right in first_meetings)
	counts = Counter(player.groupname for player in players if player.groupname)
	minimum = max(0, max(counts.values(), default=0) - (len(players) - len(first_meetings)))
	assert conflicts == minimum, (len(players), seed, counts, conflicts, minimum)
	return pairs


def memory_filesystem(module):
	"""実際のWorkbookを使い、保存先だけをメモリへ置き換える。"""
	files = {}

	def bind_save(book):
		def save(path):
			buffer = io.BytesIO()
			openpyxl.Workbook.save(book, buffer)
			files[path] = buffer.getvalue()
		book.save = save
		return book

	def remove(path):
		if path not in files:
			raise FileNotFoundError(path)
		del files[path]

	module.os = types.SimpleNamespace(remove=remove)
	module.Workbook = lambda: bind_save(openpyxl.Workbook())
	module.load_workbook = lambda path: bind_save(openpyxl.load_workbook(io.BytesIO(files[path])))
	return files


def xml_parts(data):
	with zipfile.ZipFile(io.BytesIO(data)) as archive:
		parts = {}
		for name in archive.namelist():
			root = ET.fromstring(archive.read(name))
			if name == "docProps/core.xml":
				for child in list(root):
					if child.tag.rsplit("}", 1)[-1] in ("created", "modified"):
						root.remove(child)
			parts[name] = ET.tostring(root)
		return parts


def convert_rounds(rounds, module):
	def slot(value):
		if value is None:
			return None
		if hasattr(value, "match_index"):
			return module.Ref(value.match_index)
		return module.Participant(value.name, value.kana, value.groupname)

	return [
		[module.Match(slot(match.left), slot(match.right), slot(match.winner)) for match in matches]
		for matches in rounds
	]


def check_rounds(module, players):
	before = list(players)
	rounds = module.build_full_bracket(players, seed=123)
	assert players == before
	assert rounds == module.build_full_bracket(players, seed=123)
	match_count = 0
	for round_index, matches in enumerate(rounds):
		assert len(matches) == len(rounds[0]) >> round_index
		for index, match in enumerate(matches):
			if not (module._is_bye(match.left) or module._is_bye(match.right)):
				match_count += 1
			if round_index == 0:
				assert match.winner == (match.left if module._is_bye(match.right) else None)
			else:
				for offset, value in enumerate((match.left, match.right)):
					previous = rounds[round_index - 1][2 * index + offset]
					assert value == (previous.winner or module.Ref(2 * index + offset))
	assert match_count == len(players) - 1
	assert len(rounds[-1]) == 1
	return rounds


def check_excel(old, new, old_files, new_files, rounds):
	# 抽選結果の一致は求めず、同じ対戦表を渡したときの出力形式を比較する。
	for category, init, hide in (("個人の部", True, False), ("団体の部", False, True), ("個人の部", False, False)):
		for module, bracket in ((old, convert_rounds(rounds, old)), (new, rounds)):
			with contextlib.redirect_stdout(io.StringIO()):
				module.save_bracket_xlsx(
					bracket, path="memory.xlsx", init_workbook=init,
					hide_groupname=hide, summary_name=category, match_name=category,
					match_date="2026.10.19", match_place1="第一試合場", match_place2="第二試合場",
					playername_formatter=(lambda name, category: name + "選手") if hide else None,
					groupname_formatter=None if hide else (lambda group, category: group + "団体"),
				)
		assert xml_parts(old_files["memory.xlsx"]) == xml_parts(new_files["memory.xlsx"])


def check_entrypoint(module, files, event):
	"""05_CreateTournament.pyのCreate関数を、架空TSVとメモリ保存で実行する。"""
	import ast

	tree = ast.parse((ROOT / event / "05_CreateTournament.py").read_text(encoding="utf-8"))
	function = next(node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name == "Create")
	if event == "春の大会":
		rows = [
			["Aチーム", "", "氏名"], ["", "先鋒", "選手A"], ["", "", ""],
			["Bチーム", "", "氏名"], ["", "先鋒", "選手B"], ["", "", ""],
		]
	else:
		rows = [["", "氏名", "ふりがな"], ["", "選手A", "せんしゅあ"], ["", "選手B", "せんしゅび"]]
	defines = types.SimpleNamespace(
		l_groupname=["団体A", "団体B"], download_folder="dummy", tournament_folder="dummy",
		tsv_get_all_values=lambda path: [row + [""] * (10 - len(row)) for row in rows],
		pad_list=lambda row, length: row + [""] * max(0, length - len(row)),
		get_formal_groupname=lambda name, category: name,
		get_simple_groupname=lambda name, category: name,
	)
	namespace = {"Defines": defines, "Tournament": module}
	exec(compile(ast.Module(body=[function], type_ignores=[]), "<Create>", "exec"), namespace)
	with contextlib.redirect_stdout(io.StringIO()):
		namespace["Create"]("架空の部", seed=123, init_workbook=True)
	key = "dummy\\Tournament_" + ("春" if event == "春の大会" else "秋") + ".xlsx"
	book = openpyxl.load_workbook(io.BytesIO(files[key]))
	sheet = book["架空の部"]
	assert [sheet.cell(6 + index * 2, 2).value for index in range(4)] == [1, 2, 3, 4]


def main():
	for event in ("春の大会", "秋の大会"):
		content = (ROOT / event / "Tournament.py").read_bytes()
		assert not content.startswith(b"\xef\xbb\xbf")
		assert b"\n" not in content.replace(b"\r\n", b"")
	assert (ROOT / "春の大会/Tournament.py").read_bytes() == (ROOT / "秋の大会/Tournament.py").read_bytes()
	new = load_module("tournament_under_test", ROOT / "秋の大会/Tournament.py")
	old = load_module("tournament_original", ROOT / "テスト/TournamentOld.py")

	exhaustive_count = 0
	for count in range(2, 25):
		for sizes in partitions(count):
			groups = [str(group) for group, size in enumerate(sizes) for _ in range(size)]
			players = make_players(new, groups)
			for seed in (0, 123):
				check_draw(new, players, seed)
				exhaustive_count += 1
	print(f"All group-size partitions for 2..24 players: {exhaustive_count} draws passed", flush=True)

	rng = random.Random(20261019)
	for case in range(1000):
		count = rng.randrange(2, 513)
		group_count = rng.randrange(1, min(count, 20) + 1)
		groups = [rng.choice([None, "", *map(str, range(group_count))]) for _ in range(count)]
		players = make_players(new, groups)
		seed = rng.getrandbits(64)
		pairs = check_draw(new, players, seed)
		if case % 25 == 0:
			assert pairs == new.make_first_round_pairs(players, seed)
	print("Random inputs including ungrouped players: 1000 draws passed", flush=True)

	# 不戦勝者を選べる場面では、同じ8枠内の同団体への偏りを避ける。
	for groups in (["A", "A", "B", "B", "C"], ["A", "A", "B", "B", "C", "C"]):
		players = make_players(new, groups)
		for seed in range(20):
			pairs = new.make_first_round_pairs(players, seed)
			bye_groups = [left.groupname for left, right in pairs if right.name == "BYE"]
			assert len(bye_groups) == len(set(bye_groups))
	print("Bye selection for 5 and 6 players: 40 draws passed", flush=True)

	for count in (0, 1):
		try:
			new.make_first_round_pairs(make_players(new, ["A"] * count), seed=0)
		except ValueError:
			pass
		else:
			raise AssertionError("Fewer than two players must be rejected")
	players = make_players(new, ["A", "A", "B", "B", "C", "C"])
	assert len({str(new.make_first_round_pairs(players, seed)) for seed in range(20)}) > 1

	old_files, new_files = memory_filesystem(old), memory_filesystem(new)
	for count in (2, 3, 4, 5, 6, 7, 8, 9, 16, 17, 31, 32, 33, 65, 128, 257):
		players = make_players(new, [None if index % 7 == 0 else str(index % 5) for index in range(count)])
		rounds = check_rounds(new, players)
		check_excel(old, new, old_files, new_files, rounds)
	print("Round references and Excel XML: 16 inputs x 3 export modes passed", flush=True)
	for event in ("春の大会", "秋の大会"):
		check_entrypoint(new, new_files, event)
	print("Both existing Create entrypoints passed with synthetic TSV", flush=True)


if __name__ == "__main__":
	main()
